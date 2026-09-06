from io import BytesIO

import pytest
from docx import Document

from offerlens_agent.services.documents import ResumeExtractionError, extract_resume_text


def test_extract_utf8_text_resume() -> None:
    text = extract_resume_text("resume.md", "张三\n使用 Python 和 LangGraph 完成 Agent 项目。".encode())
    assert "LangGraph" in text


def test_extract_docx_resume() -> None:
    document = Document()
    document.add_heading("张三", level=1)
    document.add_paragraph("使用 FastAPI 构建了可测试的 Agent 服务。")
    content = BytesIO()
    document.save(content)

    assert "FastAPI" in extract_resume_text("resume.docx", content.getvalue())


def test_reject_unsupported_resume_type() -> None:
    with pytest.raises(ResumeExtractionError, match="unsupported"):
        extract_resume_text("resume.exe", b"not a resume but definitely longer than twenty bytes")


def test_reject_oversized_resume_before_parsing() -> None:
    with pytest.raises(ResumeExtractionError, match="exceeds"):
        extract_resume_text("resume.txt", b"a" * 101, max_bytes=100)
