"""Safe extraction of resume text from supported document formats."""

from __future__ import annotations

from io import BytesIO
from pathlib import Path

from docx import Document
from pypdf import PdfReader

SUPPORTED_RESUME_TYPES = {".txt", ".md", ".pdf", ".docx"}


class ResumeExtractionError(ValueError):
    """Raised when a resume cannot be safely converted into text."""


def extract_resume_text(filename: str, content: bytes, max_bytes: int = 5 * 1024 * 1024) -> str:
    """Extract text while enforcing type, size, and minimum-content constraints."""
    suffix = Path(filename).suffix.lower()
    if suffix not in SUPPORTED_RESUME_TYPES:
        supported = ", ".join(sorted(SUPPORTED_RESUME_TYPES))
        raise ResumeExtractionError(f"unsupported resume type {suffix!r}; expected one of: {supported}")
    if not content:
        raise ResumeExtractionError("resume file is empty")
    if len(content) > max_bytes:
        raise ResumeExtractionError(f"resume exceeds the {max_bytes} byte limit")

    try:
        if suffix in {".txt", ".md"}:
            text = content.decode("utf-8-sig")
        elif suffix == ".pdf":
            reader = PdfReader(BytesIO(content))
            text = "\n".join(page.extract_text() or "" for page in reader.pages)
        else:
            document = Document(BytesIO(content))
            paragraphs = [paragraph.text for paragraph in document.paragraphs]
            table_cells = [cell.text for table in document.tables for row in table.rows for cell in row.cells]
            text = "\n".join([*paragraphs, *table_cells])
    except Exception as exc:  # libraries expose many format-specific exception types
        raise ResumeExtractionError(f"could not read resume: {exc}") from exc

    normalized = text.strip()
    if len(normalized) < 20:
        raise ResumeExtractionError("resume contains too little extractable text")
    return normalized
